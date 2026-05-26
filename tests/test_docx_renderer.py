"""Tests for wireghost.reports.docx_renderer — professional VA DOCX report generator."""

from __future__ import annotations

from datetime import datetime

from wireghost.models.finding import Finding
from wireghost.models.report import ScanReport
from wireghost.models.scan import Host, Port, Screenshot, Service
from wireghost.models.severity import Severity
from wireghost.reports.docx_renderer import (
    DocxRenderer,
    _group_findings_by_host,
    _group_hosts_by_subnet,
    _ip_key,
)


# ── _ip_key ─────────────────────────────────────────────────────────────────

class TestIpKey:
    def test_sorts_numerically(self):
        ips = ["10.0.0.50", "10.0.0.5", "192.168.1.1", "10.0.0.10"]
        assert sorted(ips, key=_ip_key) == [
            "10.0.0.5", "10.0.0.10", "10.0.0.50", "192.168.1.1",
        ]

    def test_different_octet_lengths(self):
        ips = ["10.0.0.100", "10.0.0.5", "172.16.254.1", "172.16.3.1"]
        assert sorted(ips, key=_ip_key) == [
            "10.0.0.5", "10.0.0.100", "172.16.3.1", "172.16.254.1",
        ]


# ── _group_hosts_by_subnet ──────────────────────────────────────────────────

def _make_host(ip, hostname="", ports=None, screenshots=None, techs=None):
    return Host(
        ip=ip, hostname=hostname, ports=ports or [],
        screenshots=screenshots or [], technologies=techs or [],
    )


class TestGroupHostsBySubnet:
    def test_groups_by_class_c(self):
        hosts = [
            _make_host("10.0.0.1"),
            _make_host("10.0.0.5"),
            _make_host("192.168.1.10"),
        ]
        result = _group_hosts_by_subnet(hosts)
        assert list(result.keys()) == ["10.0.0.0/24", "192.168.1.0/24"]
        assert [h.ip for h in result["10.0.0.0/24"]] == ["10.0.0.1", "10.0.0.5"]

    def test_sorts_subnets_by_ip(self):
        hosts = [_make_host("192.168.1.1"), _make_host("10.0.0.1")]
        result = _group_hosts_by_subnet(hosts)
        assert list(result.keys()) == ["10.0.0.0/24", "192.168.1.0/24"]


# ── _group_findings_by_host ─────────────────────────────────────────────────

def _make_finding(host, severity, title="Test"):
    return Finding(
        source="test", host=host, port="80", protocol="tcp",
        severity=severity, title=title, description="desc",
    )


class TestGroupFindingsByHost:
    def test_groups_by_host(self):
        f1 = _make_finding("10.0.0.5", Severity.HIGH, "SMB issue")
        f2 = _make_finding("10.0.0.1", Severity.CRITICAL, "SSH issue")
        result = _group_findings_by_host([f1, f2])
        assert list(result.keys()) == ["10.0.0.1", "10.0.0.5"]

    def test_sorts_by_severity_within_host(self):
        f1 = _make_finding("10.0.0.1", Severity.LOW, "Low")
        f2 = _make_finding("10.0.0.1", Severity.CRITICAL, "Critical")
        f3 = _make_finding("10.0.0.1", Severity.MEDIUM, "Medium")
        result = _group_findings_by_host([f1, f2, f3])
        severities = [f.severity for f in result["10.0.0.1"]]
        assert severities == [Severity.CRITICAL, Severity.MEDIUM, Severity.LOW]


# ── DocxRenderer._resolve_logos ─────────────────────────────────────────────

class FakeConfig:
    logo_path = ""
    prepared_by = "Test Team"
    reviewed_by = ""
    approved_by = ""


class TestResolveLogos:
    def test_empty_logo_uses_builtin_defaults(self):
        cfg = FakeConfig()
        cfg.logo_path = ""
        logo, header = DocxRenderer._resolve_logos(cfg)
        # Falls back to built-in assets if present; None if not bundled
        if logo is not None:
            assert logo.name == "logo.png"
        if header is not None:
            assert header.name == "logo_header.png"

    def test_custom_logo_that_exists(self, tmp_path):
        custom = tmp_path / "custom.png"
        custom.write_text("fake")
        cfg = FakeConfig()
        cfg.logo_path = str(custom)
        logo, _header = DocxRenderer._resolve_logos(cfg)
        assert logo == custom


# ── DocxRenderer.render — integration tests ─────────────────────────────────

def _make_port(num, proto="tcp", state="open", svc_name="", product="", version=""):
    svc = Service(name=svc_name, product=product, version=version) if svc_name else None
    return Port(number=num, protocol=proto, state=state, service=svc, service_source="nmap")


class TestRenderStructure:
    """End-to-end render tests that verify the DOCX output structure."""

    def test_minimal_report_has_expected_sections(self, tmp_path):
        report = ScanReport(
            target="test.local",
            hosts=[_make_host("10.0.0.1")],
            findings=[],
            scan_start=datetime(2026, 1, 1, 8, 0),
            scan_end=datetime(2026, 1, 1, 12, 0),
        )
        renderer = DocxRenderer()
        path = renderer.render(report, FakeConfig(), tmp_path)
        assert path.name == "security_report.docx"
        assert path.stat().st_size > 1000

        from docx import Document
        doc = Document(str(path))
        headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
        assert "Executive Summary" in headings
        assert "Target Subnets" in headings
        assert any("Identified Live Hosts" in h for h in headings)
        assert any("Open Ports" in h for h in headings)
        assert "END OF DOCUMENT" in [p.text for p in doc.paragraphs]

    def test_no_findings_skips_identified_issues(self, tmp_path):
        report = ScanReport(
            target="test",
            hosts=[_make_host("10.0.0.1")],
            findings=[],
            scan_start=datetime(2026, 1, 1),
            scan_end=datetime(2026, 1, 1),
        )
        renderer = DocxRenderer()
        path = renderer.render(report, FakeConfig(), tmp_path)
        from docx import Document
        doc = Document(str(path))
        headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
        assert not any("Identified issues" in h for h in headings)

    def test_section_3_is_live_hosts(self, tmp_path):
        report = ScanReport(
            target="test",
            hosts=[_make_host("10.0.0.1")],
            findings=[],
            scan_start=datetime(2026, 1, 1),
            scan_end=datetime(2026, 1, 1),
        )
        renderer = DocxRenderer()
        path = renderer.render(report, FakeConfig(), tmp_path)
        from docx import Document
        doc = Document(str(path))
        h1s = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
        live_hosts = [h for h in h1s if "Live Hosts" in h]
        assert len(live_hosts) == 1
        assert live_hosts[0].startswith("3.")


class TestRenderSeverityColoring:
    """Verify severity colors appear in headings and table."""

    def test_finding_headings_are_colored(self, tmp_path):
        report = ScanReport(
            target="test",
            hosts=[_make_host("10.0.0.1")],
            findings=[
                _make_finding("10.0.0.1", Severity.CRITICAL, "Critical vuln"),
                _make_finding("10.0.0.1", Severity.INFO, "Info note"),
            ],
            scan_start=datetime(2026, 1, 1),
            scan_end=datetime(2026, 1, 1),
        )
        renderer = DocxRenderer()
        path = renderer.render(report, FakeConfig(), tmp_path)
        from docx import Document
        from wireghost.reports.docx_renderer import _GREEN, _SEVERITY_RGB

        doc = Document(str(path))
        for para in doc.paragraphs:
            if para.style.name == "Heading 2" and "[" in para.text:
                runs = para.runs
                assert len(runs) >= 3
                # First and last runs green, middle run severity-colored
                assert runs[0].font.color.rgb == _GREEN
                assert runs[2].font.color.rgb == _GREEN
                sev_label = runs[1].text.strip()
                expected_rgb = _SEVERITY_RGB.get(
                    Severity[sev_label] if sev_label in ("CRITICAL", "HIGH", "MEDIUM", "LOW", "INFO") else None
                )
                if expected_rgb:
                    assert runs[1].font.color.rgb == expected_rgb

    def test_severity_table_has_colored_labels(self, tmp_path):
        report = ScanReport(
            target="test",
            hosts=[_make_host("10.0.0.1")],
            findings=[
                _make_finding("10.0.0.1", Severity.CRITICAL, "vuln"),
                _make_finding("10.0.0.1", Severity.HIGH, "vuln"),
            ],
            scan_start=datetime(2026, 1, 1),
            scan_end=datetime(2026, 1, 1),
        )
        renderer = DocxRenderer()
        path = renderer.render(report, FakeConfig(), tmp_path)
        from docx import Document
        from wireghost.reports.docx_renderer import _SEVERITY_RGB

        doc = Document(str(path))
        colored = 0
        for table in doc.tables:
            for row in table.rows:
                cell_text = row.cells[0].text.strip()
                for sev, rgb in _SEVERITY_RGB.items():
                    if sev.name.capitalize() == cell_text:
                        run = row.cells[0].paragraphs[0].runs[0]
                        assert run.font.color.rgb == rgb
                        colored += 1
        assert colored >= 2  # CRITICAL + HIGH


class TestRenderScreenshots:
    """Verify screenshot positioning — inline vs orphan."""

    def test_screenshots_inline_when_host_has_findings(self, tmp_path):
        ss_dir = tmp_path.parent / "ips" / "10.0.0.1" / "web" / "screenshots"
        ss_dir.mkdir(parents=True)
        (ss_dir / "shot.png").write_bytes(
            b"\x89PNG\r\n\x1a\n" + b"\x00" * 100  # minimal valid PNG header
        )

        report = ScanReport(
            target="test",
            hosts=[
                Host(
                    ip="10.0.0.1", hostname="web",
                    ports=[_make_port(80, svc_name="http", product="nginx")],
                    screenshots=[Screenshot(url="http://10.0.0.1", filename="shot.png", title="Home")],
                ),
            ],
            findings=[
                _make_finding("10.0.0.1", Severity.HIGH, "nginx vuln"),
            ],
            scan_start=datetime(2026, 1, 1),
            scan_end=datetime(2026, 1, 1),
        )
        renderer = DocxRenderer()
        path = renderer.render(report, FakeConfig(), tmp_path)
        from docx import Document
        doc = Document(str(path))

        # Screenshot should be INLINE (inside identified issues), not standalone
        inline_captions = [p.text for p in doc.paragraphs if "Web Screenshot:" in p.text]
        assert len(inline_captions) == 1
        # No standalone screenshots heading
        h1s = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
        assert not any("Web Screenshots" in h for h in h1s)

    def test_standalone_screenshots_for_orphan_hosts(self, tmp_path):
        ss_dir = tmp_path.parent / "ips" / "10.0.0.2" / "web" / "screenshots"
        ss_dir.mkdir(parents=True)
        (ss_dir / "shot.png").write_bytes(
            b"\x89PNG\r\n\x1a\n" + b"\x00" * 100
        )

        report = ScanReport(
            target="test",
            hosts=[
                _make_host("10.0.0.1"),  # has finding
                Host(
                    ip="10.0.0.2", hostname="web",
                    ports=[_make_port(80, svc_name="http")],
                    screenshots=[Screenshot(url="http://10.0.0.2", filename="shot.png")],
                ),
            ],
            findings=[
                _make_finding("10.0.0.1", Severity.LOW, "ssh weak"),
            ],
            scan_start=datetime(2026, 1, 1),
            scan_end=datetime(2026, 1, 1),
        )
        renderer = DocxRenderer()
        path = renderer.render(report, FakeConfig(), tmp_path)
        from docx import Document
        doc = Document(str(path))

        # 10.0.0.2 has screenshots but no findings → standalone section
        h1s = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
        assert any("Web Screenshots" in h for h in h1s)
        # No inline captions (10.0.0.1 has findings but no screenshots)
        assert not any("Web Screenshot:" in p.text for p in doc.paragraphs)

    def test_no_screenshots_section_when_none_exist(self, tmp_path):
        report = ScanReport(
            target="test",
            hosts=[_make_host("10.0.0.1")],
            findings=[_make_finding("10.0.0.1", Severity.LOW, "test")],
            scan_start=datetime(2026, 1, 1),
            scan_end=datetime(2026, 1, 1),
        )
        renderer = DocxRenderer()
        path = renderer.render(report, FakeConfig(), tmp_path)
        from docx import Document
        doc = Document(str(path))
        headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
        assert not any("Web Screenshots" in h for h in headings)


class TestRenderSectionNumbering:
    def test_executive_summary_unnumbered(self, tmp_path):
        report = ScanReport(
            target="test", hosts=[_make_host("10.0.0.1")], findings=[],
            scan_start=datetime(2026, 1, 1), scan_end=datetime(2026, 1, 1),
        )
        renderer = DocxRenderer()
        path = renderer.render(report, FakeConfig(), tmp_path)
        from docx import Document
        doc = Document(str(path))
        h1s = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
        assert "Executive Summary" in h1s[0] and not h1s[0][0].isdigit()
        assert "Target Subnets" in h1s[1] and not h1s[1][0].isdigit()
